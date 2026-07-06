"""FaberLoom SL2: knowledge-base ingestion, retrieval and draft helpers.

KB design for SL2:
- Text sources (.md/.txt/.pdf) are chunked into ~900 char segments and indexed with FTS5.
- Tabular sources (.csv/.xlsx and tables inside PDFs) are parsed into kb_fact rows
  for source-to-field checking.
- Retrieval is workspace-scoped and keyword-based (no embeddings yet).
- Draft generation consumes an evidence pack and must cite real KB data.
"""

from __future__ import annotations

import csv
import difflib
import io
import json
import re
import sqlite3
from datetime import datetime, timezone
from typing import Any

import openpyxl
import pdfplumber

from .context import Context, enforce_tenant_scoped
from .db import new_id, record_editorial_event, row_to_dict, transaction, utc_now, workspace_seal_id
from .kb_extractors import extract_document
from .models import SCHEMA_VERSION
from .seal import (
    assert_workspace_hmac_for_row,
    compute_workspace_hmac_for_row,
    verify_workspace_hmac_for_row,
)
from .security import assert_safe_kb_source


PARSER_VERSIONS = {
    "md": "builtin",
    "txt": "builtin",
    "csv": "builtin",
    "xlsx": f"openpyxl {openpyxl.__version__}",
    "pdf": f"pdfplumber {pdfplumber.__version__}",
}


KB_CHUNK_SIZE = 900
KB_CHUNK_OVERLAP = 100


def _today() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _slugify_locator(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()[:120]


def _chunk_text(text: str, size: int = KB_CHUNK_SIZE, overlap: int = KB_CHUNK_OVERLAP) -> list[str]:
    """Produce overlapping chunks of plain text."""

    text = text.strip()
    if not text:
        return []
    if len(text) <= size:
        return [text]

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = start + size
        if end >= len(text):
            chunks.append(text[start:].strip())
            break
        # Try to break at a newline or space near the boundary.
        break_at = text.rfind("\n", end - 40, end)
        if break_at == -1:
            break_at = text.rfind(" ", end - 40, end)
        if break_at == -1 or break_at <= start:
            break_at = end
        chunks.append(text[start:break_at].strip())
        start = max(start + 1, break_at - overlap)
    return chunks


def _workspace_kb_scope(
    ctx: Context,
    conn: sqlite3.Connection,
) -> list[str]:
    """Return the workspace ids that contribute KB data to the current context.

    If the current workspace inherits KB from its parent and both share the
    same tenant, the parent's workspace_id is included in the scope.
    """

    workspace_id = ctx.require_scoped_workspace()
    row = conn.execute(
        "SELECT parent_id, inherits_kb, tenant_id FROM workspace WHERE id = ?",
        (workspace_id,),
    ).fetchone()
    if row and row["inherits_kb"] and row["parent_id"] and row["tenant_id"] == ctx.tenant_id:
        return [workspace_id, row["parent_id"]]
    return [workspace_id]


def _workspace_seals(
    ctx: Context,
    conn: sqlite3.Connection,
    scope: list[str],
) -> dict[str, str]:
    """Return seal_id per workspace for HMAC verification inside a KB scope."""

    if not scope:
        return {}
    placeholders = ",".join("?" for _ in scope)
    rows = conn.execute(
        f"SELECT id, seal_id FROM workspace WHERE id IN ({placeholders}) AND tenant_id = ?",
        (*scope, ctx.tenant_id),
    ).fetchall()
    return {row["id"]: row["seal_id"] for row in rows if row["seal_id"]}


def _is_date_stale(valid_from: str | None, valid_until: str | None) -> bool:
    """Return True if a validity window is currently stale or not yet valid."""

    today = datetime.fromisoformat(_today().replace("Z", "+00:00")).date()
    if valid_from:
        try:
            from_dt = datetime.fromisoformat(valid_from.replace("Z", "+00:00")).date()
            if from_dt > today:
                return True
        except ValueError:
            pass
    if valid_until:
        try:
            until_dt = datetime.fromisoformat(valid_until.replace("Z", "+00:00")).date()
            if until_dt < today:
                return True
        except ValueError:
            pass
    return False


def _table_rows_to_facts(
    ctx: Context,
    conn: sqlite3.Connection,
    *,
    source_id: str,
    table_name: str | None,
    source_sheet: str | None,
    rows: list[dict[str, str]],
    source_version: str,
    extraction_method: str,
    seal_id: str | None = None,
    extraction_warnings: list[str] | None = None,
) -> None:
    """Insert table rows as kb_fact records.

    Heuristic: the first column is treated as entity_key. Columns named
    vigente_desde/hasta/valid_from/valid_until are used for validity windows;
    moneda/currency is propagated to other facts.

    Rows with stale or not-yet-valid windows are inserted (so the data is
    preserved) but flagged in ``extraction_warnings``.
    """

    if not rows:
        return

    headers = list(rows[0].keys())
    entity_header = headers[0]
    valid_from_header = next(
        (h for h in headers if h.lower() in {"vigente_desde", "valid_from", "desde"}),
        None,
    )
    valid_until_header = next(
        (h for h in headers if h.lower() in {"vigente_hasta", "valid_until", "hasta"}),
        None,
    )
    currency_header = next(
        (h for h in headers if h.lower() in {"moneda", "currency"}),
        None,
    )

    locator_prefix = f"{table_name}: " if table_name else ""
    workspace_id = ctx.require_scoped_workspace()
    tenant_id = ctx.tenant_id
    if seal_id is None:
        seal_id = workspace_seal_id(ctx, conn)

    for row_index, row in enumerate(rows, start=2):  # row 1 is header
        entity_key = row.get(entity_header, "").strip()
        if not entity_key:
            continue

        valid_from = row.get(valid_from_header, "").strip() if valid_from_header else None
        valid_until = row.get(valid_until_header, "").strip() if valid_until_header else None
        currency = row.get(currency_header, "").strip() if currency_header else None

        if _is_date_stale(valid_from, valid_until) and extraction_warnings is not None:
            extraction_warnings.append(
                f"Row {row_index} ({entity_key}) has a stale or not-yet-valid date window."
            )

        for header in headers:
            value = row.get(header, "").strip()
            if not value or header in {valid_from_header, valid_until_header, currency_header}:
                continue
            if header == entity_header:
                continue

            fact_id = new_id("fact")
            fact_hmac_row = {
                "id": fact_id,
                "workspace_id": workspace_id,
                "source_id": source_id,
                "entity_key": entity_key,
                "field_name": header,
                "field_value": value,
                "source_version": source_version,
            }
            hmac = compute_workspace_hmac_for_row(seal_id, fact_hmac_row, "kb_fact")
            conn.execute(
                """
                INSERT INTO kb_fact (
                    id, workspace_id, source_id, entity_key, field_name, field_value,
                    currency, valid_from, valid_until, source_locator, source_version,
                    extraction_method, source_sheet, schema_version, workspace_hmac,
                    tenant_id, created_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    fact_id,
                    workspace_id,
                    source_id,
                    entity_key,
                    header,
                    value,
                    currency,
                    valid_from or None,
                    valid_until or None,
                    f"{locator_prefix}row {row_index}",
                    source_version,
                    extraction_method,
                    source_sheet,
                    SCHEMA_VERSION,
                    hmac,
                    tenant_id,
                    utc_now(),
                ),
            )


def ingest_kb_source(
    ctx: Context,
    conn: sqlite3.Connection,
    *,
    title: str,
    source_type: str,
    content_text: str | None = None,
    content_blob: bytes | None = None,
    source_version: str = "v1",
    level: int = 0,
    approved_by: str | None = None,
    file_name: str | None = None,
    mime_type: str | None = None,
    file_size: int | None = None,
) -> dict[str, Any]:
    """Ingest a KB source and materialize chunks/facts.

    Supports text sources (md/txt), tabular sources (csv/xlsx), and PDFs.
    Either *content_text* (for text/csv) or *content_blob* (for binary files)
    must be provided.
    """

    enforce_tenant_scoped(ctx)

    if source_type not in {"md", "txt", "csv", "xlsx", "pdf"}:
        raise ValueError(f"Unsupported KB source type: {source_type}")

    workspace_id = ctx.require_scoped_workspace()
    source_id = new_id("kbs")
    now = utc_now()

    extracted_text = ""
    extracted_tables: list[dict[str, Any]] | None = None
    extraction_warnings: list[str] | None = None
    extraction_method = source_type
    parser_version = PARSER_VERSIONS.get(source_type, "unknown")

    if source_type in {"md", "txt", "csv"}:
        if content_text is None:
            if content_blob is None:
                raise ValueError(f"content_text or content_blob required for source type {source_type}")
            try:
                content_text = content_blob.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise ValueError(f"Could not decode {source_type} file as UTF-8") from exc
        assert_safe_kb_source(source_type, content_text)
        extracted_text = content_text
        if source_type == "csv":
            try:
                reader = csv.DictReader(io.StringIO(content_text))
                extracted_tables = [{"name": None, "rows": list(reader)}]
            except Exception as exc:
                raise ValueError(f"Invalid CSV: {exc}") from exc
            extraction_warnings = []
    elif source_type in {"xlsx", "pdf"}:
        if content_blob is None:
            raise ValueError(f"content_blob required for source type {source_type}")
        doc = extract_document(
            content_blob,
            file_name=file_name,
            mime_type=mime_type,
        )
        extracted_text = doc.text
        extraction_warnings = doc.warnings or []
        extracted_tables = [
            {"name": table.name, "rows": table.rows}
            for table in (doc.tables or [])
            if table.rows
        ]
        # Re-sanitize extracted text against known injection vectors.
        assert_safe_kb_source("txt", extracted_text)
        extraction_method = f"{source_type}_extractor"
        if file_size is None:
            file_size = len(content_blob)
    else:
        raise ValueError(f"Unsupported KB source type: {source_type}")

    meta: dict[str, Any] = {
        "title": title,
        "type": source_type,
        "source_version": source_version,
        "extraction_method": extraction_method,
        "parser_version": parser_version,
    }
    with transaction(conn):
        conn.execute(
            f"""
            INSERT INTO kb_source (
                id, workspace_id, tenant_id, user_id, actor_id, actor_role_at_decision,
                type, title, content_text, content_blob, meta_json, source_version, schema_version,
                level, file_name, mime_type, file_size, parser_version, approved_by, workspace_hmac,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                source_id,
                workspace_id,
                ctx.tenant_id,
                ctx.user_id,
                ctx.resolved_actor_id(),
                ctx.actor_role_at_decision,
                source_type,
                title.strip(),
                extracted_text or None,
                content_blob,
                json.dumps(meta, ensure_ascii=False, sort_keys=True),
                source_version,
                SCHEMA_VERSION,
                level,
                file_name,
                mime_type,
                file_size,
                parser_version,
                approved_by,
                None,  # workspace_hmac set below after seal lookup
                now,
            ),
        )
        seal_id = workspace_seal_id(ctx, conn)
        source_hmac = compute_workspace_hmac_for_row(
            seal_id,
            {
                "id": source_id,
                "workspace_id": workspace_id,
                "type": source_type,
                "title": title.strip(),
                "source_version": source_version,
            },
            "kb_source",
        )
        conn.execute(
            "UPDATE kb_source SET workspace_hmac = ? WHERE id = ?",
            (source_hmac, source_id),
        )

        if extracted_text:
            for index, chunk in enumerate(_chunk_text(extracted_text)):
                conn.execute(
                    """
                    INSERT INTO kb_chunk (
                        id, workspace_id, source_id, chunk_index, content_text,
                        source_locator, source_version, schema_version, tenant_id, created_at
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        new_id("chunk"),
                        workspace_id,
                        source_id,
                        index,
                        chunk,
                        f"chunk {index}",
                        source_version,
                        SCHEMA_VERSION,
                        ctx.tenant_id,
                        now,
                    ),
                )

        if extracted_tables:
            for table in extracted_tables:
                _table_rows_to_facts(
                    ctx,
                    conn,
                    source_id=source_id,
                    table_name=table.get("name"),
                    source_sheet=table.get("name"),
                    rows=table["rows"],
                    source_version=source_version,
                    extraction_method=extraction_method,
                    seal_id=seal_id,
                    extraction_warnings=extraction_warnings,
                )

        if extraction_warnings:
            meta["extraction_warnings"] = extraction_warnings
            if any("stale or not-yet-valid date window" in w for w in extraction_warnings):
                meta["stale_data_block"] = True
            conn.execute(
                "UPDATE kb_source SET meta_json = ? WHERE id = ?",
                (json.dumps(meta, ensure_ascii=False, sort_keys=True), source_id),
            )

    row = conn.execute(
        """
        SELECT id, workspace_id, tenant_id, user_id, actor_id, actor_role_at_decision,
               type, title, content_text, meta_json, source_version, schema_version,
               file_name, mime_type, file_size, parser_version, approved_by,
               workspace_hmac, created_at
        FROM kb_source
        WHERE id = ? AND workspace_id = ? AND tenant_id = ?
        """,
        (source_id, workspace_id, ctx.tenant_id),
    ).fetchone()
    assert row is not None
    return row_to_dict(row)


def list_kb_sources(ctx: Context, conn: sqlite3.Connection) -> list[dict[str, Any]]:
    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    rows = conn.execute(
        f"""
        SELECT id, workspace_id, tenant_id, user_id, actor_id, actor_role_at_decision,
               type, title, content_text, meta_json, source_version, schema_version,
               file_name, mime_type, file_size, parser_version, approved_by,
               workspace_hmac, created_at
        FROM kb_source
        WHERE workspace_id IN ({placeholders}) AND tenant_id = ?
        ORDER BY created_at DESC
        """,
        (*scope, ctx.tenant_id),
    ).fetchall()
    seals = _workspace_seals(ctx, conn, scope)
    results: list[dict[str, Any]] = []
    for row in rows:
        seal_id = seals.get(row["workspace_id"])
        if seal_id and verify_workspace_hmac_for_row(seal_id, row_to_dict(row), "kb_source"):
            results.append(row_to_dict(row))
    return results


def get_kb_source(
    ctx: Context, conn: sqlite3.Connection, source_id: str
) -> dict[str, Any] | None:
    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    row = conn.execute(
        f"""
        SELECT id, workspace_id, tenant_id, user_id, actor_id, actor_role_at_decision,
               type, title, content_text, meta_json, source_version, schema_version,
               file_name, mime_type, file_size, parser_version, approved_by,
               workspace_hmac, created_at
        FROM kb_source
        WHERE id = ? AND workspace_id IN ({placeholders}) AND tenant_id = ?
        """,
        (source_id, *scope, ctx.tenant_id),
    ).fetchone()
    if row is None:
        return None
    seals = _workspace_seals(ctx, conn, scope)
    seal_id = seals.get(row["workspace_id"])
    if not seal_id:
        return None
    assert_workspace_hmac_for_row(seal_id, row_to_dict(row), "kb_source")
    return row_to_dict(row)


def delete_kb_source(ctx: Context, conn: sqlite3.Connection, source_id: str) -> bool:
    workspace_id = ctx.require_scoped_workspace()
    with transaction(conn):
        cursor = conn.execute(
            "DELETE FROM kb_source WHERE id = ? AND workspace_id = ? AND tenant_id = ?",
            (source_id, workspace_id, ctx.tenant_id),
        )
    return cursor.rowcount > 0


def search_kb_chunks(
    ctx: Context,
    conn: sqlite3.Connection,
    query: str,
    *,
    limit: int = 5,
) -> list[dict[str, Any]]:
    """Keyword retrieval via FTS5, scoped to the current workspace and its inherited KB."""

    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    # Escape FTS5 special chars conservatively and OR the terms so any hit
    # returns evidence (we re-rank/filter downstream).
    terms = [term.strip() for term in query.split() if term.strip()]
    if not terms:
        return []
    safe_query = " OR ".join(
        f'"{term.replace('"', '""')}"'
        for term in terms
    )

    rows = conn.execute(
        f"""
        SELECT c.id, c.workspace_id, c.source_id, c.chunk_index, c.content_text,
               c.source_locator, c.source_version, c.created_at
        FROM kb_chunk_fts f
        JOIN kb_chunk c ON c.rowid = f.rowid
        JOIN kb_source s ON s.id = c.source_id
        WHERE kb_chunk_fts MATCH ?
          AND c.workspace_id IN ({placeholders})
          AND s.workspace_id IN ({placeholders})
          AND c.tenant_id = ?
          AND s.tenant_id = ?
        ORDER BY rank
        LIMIT ?
        """,
        (safe_query, *scope, *scope, ctx.tenant_id, ctx.tenant_id, limit),
    ).fetchall()
    return [row_to_dict(row) for row in rows]


def search_kb_facts(
    ctx: Context,
    conn: sqlite3.Connection,
    query: str,
    *,
    limit: int = 10,
) -> list[dict[str, Any]]:
    """Search kb_fact rows by entity_key or field_value, scoped to workspace.

    Rows with a broken workspace HMAC seal are silently discarded so a moved
    row cannot be read from another workspace.
    """

    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    pattern = f"%{query}%"
    rows = conn.execute(
        f"""
        SELECT id, workspace_id, source_id, entity_key, field_name, field_value,
               unit, currency, valid_from, valid_until, source_locator,
               source_version, extraction_method, source_sheet, workspace_hmac,
               tenant_id, created_at
        FROM kb_fact
        WHERE workspace_id IN ({placeholders})
          AND tenant_id = ?
          AND (entity_key LIKE ? OR field_value LIKE ?)
        ORDER BY created_at DESC
        LIMIT ?
        """,
        (*scope, ctx.tenant_id, pattern, pattern, limit),
    ).fetchall()
    seals = _workspace_seals(ctx, conn, scope)
    results: list[dict[str, Any]] = []
    for row in rows:
        seal_id = seals.get(row["workspace_id"])
        if seal_id and verify_workspace_hmac_for_row(seal_id, row_to_dict(row), "kb_fact"):
            results.append(row_to_dict(row))
    return results


def get_kb_fact_by_id(
    ctx: Context, conn: sqlite3.Connection, fact_id: str
) -> dict[str, Any] | None:
    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    row = conn.execute(
        f"""
        SELECT id, workspace_id, source_id, entity_key, field_name, field_value,
               unit, currency, valid_from, valid_until, source_locator,
               source_version, extraction_method, source_sheet, workspace_hmac,
               tenant_id, created_at
        FROM kb_fact
        WHERE id = ? AND workspace_id IN ({placeholders}) AND tenant_id = ?
        """,
        (fact_id, *scope, ctx.tenant_id),
    ).fetchone()
    if row is None:
        return None
    seals = _workspace_seals(ctx, conn, scope)
    seal_id = seals.get(row["workspace_id"])
    if not seal_id:
        return None
    assert_workspace_hmac_for_row(seal_id, row_to_dict(row), "kb_fact")
    return row_to_dict(row)


def get_kb_chunk_by_id(
    ctx: Context, conn: sqlite3.Connection, chunk_id: str
) -> dict[str, Any] | None:
    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    row = conn.execute(
        f"""
        SELECT id, workspace_id, source_id, chunk_index, content_text,
               source_locator, source_version, tenant_id, created_at
        FROM kb_chunk
        WHERE id = ? AND workspace_id IN ({placeholders}) AND tenant_id = ?
        """,
        (chunk_id, *scope, ctx.tenant_id),
    ).fetchone()
    return row_to_dict(row) if row else None


def find_kb_fact(
    ctx: Context,
    conn: sqlite3.Connection,
    *,
    source_id: str,
    field_name: str,
    field_value: str,
) -> dict[str, Any] | None:
    """Find a kb_fact by exact source, field and value (workspace-scoped)."""

    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    row = conn.execute(
        f"""
        SELECT id, workspace_id, source_id, entity_key, field_name, field_value,
               unit, currency, valid_from, valid_until, source_locator,
               source_version, extraction_method, source_sheet, workspace_hmac,
               tenant_id, created_at
        FROM kb_fact
        WHERE workspace_id IN ({placeholders})
          AND tenant_id = ?
          AND source_id = ? AND field_name = ? AND field_value = ?
        ORDER BY created_at DESC
        LIMIT 1
        """,
        (*scope, ctx.tenant_id, source_id, field_name, field_value),
    ).fetchone()
    if row is None:
        return None
    seals = _workspace_seals(ctx, conn, scope)
    seal_id = seals.get(row["workspace_id"])
    if not seal_id:
        return None
    assert_workspace_hmac_for_row(seal_id, row_to_dict(row), "kb_fact")
    return row_to_dict(row)


def get_kb_facts_by_source(
    ctx: Context, conn: sqlite3.Connection, source_id: str
) -> list[dict[str, Any]]:
    """Return all kb_fact rows for a source (workspace-scoped)."""

    scope = _workspace_kb_scope(ctx, conn)
    placeholders = ",".join("?" for _ in scope)
    rows = conn.execute(
        f"""
        SELECT id, workspace_id, source_id, entity_key, field_name, field_value,
               unit, currency, valid_from, valid_until, source_locator,
               source_version, extraction_method, source_sheet, workspace_hmac,
               tenant_id, created_at
        FROM kb_fact
        WHERE workspace_id IN ({placeholders}) AND tenant_id = ? AND source_id = ?
        ORDER BY entity_key, field_name
        """,
        (*scope, ctx.tenant_id, source_id),
    ).fetchall()
    seals = _workspace_seals(ctx, conn, scope)
    results: list[dict[str, Any]] = []
    for row in rows:
        seal_id = seals.get(row["workspace_id"])
        if seal_id and verify_workspace_hmac_for_row(seal_id, row_to_dict(row), "kb_fact"):
            results.append(row_to_dict(row))
    return results


# -----------------------------------------------------------------------------
# Draft helpers
# -----------------------------------------------------------------------------


def _serialize_draft(row: sqlite3.Row) -> dict[str, Any]:
    data = row_to_dict(row)
    data["requires_confirmation"] = bool(data.get("requires_confirmation", 0))
    return data


def _edit_pct(original: str | None, current: str | None) -> float | None:
    """Return edit percentage (0-100) between original and current body."""

    if original is None or current is None:
        return None
    if original == current:
        return 0.0
    # SequenceMatcher ratio is in [0, 1]; convert to percentage.
    ratio = difflib.SequenceMatcher(None, original, current).ratio()
    return round((1.0 - ratio) * 100.0, 2)


def insert_draft(
    ctx: Context,
    conn: sqlite3.Connection,
    *,
    chat_id: str | None,
    task: str,
    subject: str | None,
    body_md: str,
    hard_facts: list[dict[str, Any]],
    sources: list[dict[str, Any]],
    blockers: list[str],
    warnings: list[str],
    requires_confirmation: bool,
    status: str = "draft",
    source_version: str | None = None,
    original_body_md: str | None = None,
) -> dict[str, Any]:
    workspace_id = ctx.require_scoped_workspace()
    draft_id = new_id("draft")
    now = utc_now()
    seal_id = workspace_seal_id(ctx, conn)
    # The first persisted body is the baseline for measuring human edits.
    original_body_md = original_body_md if original_body_md is not None else body_md
    hmac = compute_workspace_hmac_for_row(
        seal_id,
        {
            "id": draft_id,
            "workspace_id": workspace_id,
            "task": task,
            "subject": subject,
            "body_md": body_md,
            "status": status,
            "source_version": source_version,
        },
        "draft",
    )

    conn.execute(
        """
        INSERT INTO draft (
            id, workspace_id, chat_id, tenant_id, user_id, actor_id,
            actor_role_at_decision, task, subject, body_md, hard_facts_json,
            sources_json, blockers_json, warnings_json, requires_confirmation,
            status, schema_version, source_version, workspace_hmac, created_at,
            updated_at, original_body_md, edit_pct
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            draft_id,
            workspace_id,
            chat_id,
            ctx.tenant_id,
            ctx.user_id,
            ctx.resolved_actor_id(),
            ctx.actor_role_at_decision,
            task,
            subject,
            body_md,
            json.dumps(hard_facts, ensure_ascii=False),
            json.dumps(sources, ensure_ascii=False),
            json.dumps(blockers, ensure_ascii=False),
            json.dumps(warnings, ensure_ascii=False),
            1 if requires_confirmation else 0,
            status,
            SCHEMA_VERSION,
            source_version,
            hmac,
            now,
            now,
            original_body_md,
            None,
        ),
    )
    row = conn.execute(
        """
        SELECT * FROM draft
        WHERE id = ? AND workspace_id = ? AND tenant_id = ?
        """,
        (draft_id, workspace_id, ctx.tenant_id),
    ).fetchone()
    assert row is not None
    return _serialize_draft(row)


def get_draft(
    ctx: Context, conn: sqlite3.Connection, draft_id: str
) -> dict[str, Any] | None:
    workspace_id = ctx.require_scoped_workspace()
    row = conn.execute(
        """
        SELECT * FROM draft
        WHERE id = ? AND workspace_id = ? AND tenant_id = ?
        """,
        (draft_id, workspace_id, ctx.tenant_id),
    ).fetchone()
    if row is None:
        return None
    seal_id = workspace_seal_id(ctx, conn)
    assert_workspace_hmac_for_row(seal_id, row_to_dict(row), "draft")
    return _serialize_draft(row)


def list_drafts(
    ctx: Context,
    conn: sqlite3.Connection,
    *,
    status: str | None = None,
    limit: int = 100,
) -> list[dict[str, Any]]:
    workspace_id = ctx.require_scoped_workspace()
    sql = """
        SELECT * FROM draft
        WHERE workspace_id = ? AND tenant_id = ?
    """
    params: list[Any] = [workspace_id, ctx.tenant_id]
    if status is not None:
        sql += " AND status = ?"
        params.append(status)
    sql += " ORDER BY created_at DESC LIMIT ?"
    params.append(limit)
    rows = conn.execute(sql, params).fetchall()
    seal_id = workspace_seal_id(ctx, conn)
    results: list[dict[str, Any]] = []
    for row in rows:
        if verify_workspace_hmac_for_row(seal_id, row_to_dict(row), "draft"):
            results.append(_serialize_draft(row))
    return results


def update_draft(
    ctx: Context,
    conn: sqlite3.Connection,
    draft_id: str,
    *,
    subject: str | None = None,
    body_md: str | None = None,
) -> dict[str, Any] | None:
    workspace_id = ctx.require_scoped_workspace()
    # Import is deferred to avoid a circular dependency with draft_engine.
    from .draft_engine import revalidate_draft

    with transaction(conn):
        row = conn.execute(
            """
            SELECT * FROM draft
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (draft_id, workspace_id, ctx.tenant_id),
        ).fetchone()
        if row is None:
            return None

        if row["status"] not in {"draft", "pending_approval"}:
            raise ValueError("Cannot edit draft after approval/rejection/export")

        updates: list[str] = []
        params: list[Any] = []
        if subject is not None:
            updates.append("subject = ?")
            params.append(subject)
        if body_md is not None:
            updates.append("body_md = ?")
            params.append(body_md)

        # Preserve the first persisted body as the edit baseline.
        current_original = row["original_body_md"]
        if current_original is None:
            updates.append("original_body_md = ?")
            params.append(row["body_md"])

        # Re-validate hard facts after any edit so "zero invented data" holds.
        new_body = body_md if body_md is not None else row["body_md"]
        reval = revalidate_draft(ctx, conn, dict(row), body_md=new_body)
        updates.append("blockers_json = ?")
        params.append(json.dumps(reval["blockers"], ensure_ascii=False))
        updates.append("warnings_json = ?")
        params.append(json.dumps(reval["warnings"], ensure_ascii=False))
        updates.append("requires_confirmation = ?")
        params.append(1 if reval["requires_confirmation"] else 0)

        seal_id = workspace_seal_id(ctx, conn)
        new_subject = subject if subject is not None else row["subject"]
        new_body = body_md if body_md is not None else row["body_md"]
        hmac = compute_workspace_hmac_for_row(
            seal_id,
            {
                "id": draft_id,
                "workspace_id": workspace_id,
                "task": row["task"],
                "subject": new_subject,
                "body_md": new_body,
                "status": row["status"],
                "source_version": row["source_version"],
            },
            "draft",
        )
        updates.append("workspace_hmac = ?")
        params.append(hmac)

        updates.append("updated_at = ?")
        params.append(utc_now())
        params.extend([draft_id, workspace_id, ctx.tenant_id])

        conn.execute(
            f"UPDATE draft SET {', '.join(updates)} WHERE id = ? AND workspace_id = ? AND tenant_id = ?",
            params,
        )

    row = conn.execute(
        """
        SELECT * FROM draft
        WHERE id = ? AND workspace_id = ? AND tenant_id = ?
        """,
        (draft_id, workspace_id, ctx.tenant_id),
    ).fetchone()
    if row is None:
        return None
    assert_workspace_hmac_for_row(seal_id, row_to_dict(row), "draft")
    return _serialize_draft(row)


def approve_draft(
    ctx: Context,
    conn: sqlite3.Connection,
    draft_id: str,
    *,
    confirmed: bool = False,
    reason: str | None = None,
    urgency: int = 0,
) -> dict[str, Any] | None:
    workspace_id = ctx.require_scoped_workspace()
    with transaction(conn):
        row = conn.execute(
            """
            SELECT * FROM draft
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (draft_id, workspace_id, ctx.tenant_id),
        ).fetchone()
        if row is None:
            return None

        data = _serialize_draft(row)
        if data["blockers_json"] != "[]":
            raise ValueError("Cannot approve draft with blockers")
        if data["requires_confirmation"] and not confirmed:
            raise ValueError(
                "Draft requires explicit confirmation before approval; retry with confirmed=true"
            )

        # Ensure the original body baseline is set and compute edit_pct.
        original_body = row["original_body_md"] if row["original_body_md"] is not None else row["body_md"]
        edit_pct = _edit_pct(original_body, row["body_md"])

        seal_id = workspace_seal_id(ctx, conn)
        hmac = compute_workspace_hmac_for_row(
            seal_id,
            {
                "id": draft_id,
                "workspace_id": workspace_id,
                "task": row["task"],
                "subject": row["subject"],
                "body_md": row["body_md"],
                "status": "approved",
                "source_version": row["source_version"],
            },
            "draft",
        )
        conn.execute(
            """
            UPDATE draft
            SET status = 'approved', approved_by = ?, approved_at = ?, updated_at = ?,
                workspace_hmac = ?, urgency = ?, reason = ?, original_body_md = ?, edit_pct = ?
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (
                ctx.resolved_actor_id(),
                utc_now(),
                utc_now(),
                hmac,
                urgency,
                reason,
                original_body,
                edit_pct,
                draft_id,
                workspace_id,
                ctx.tenant_id,
            ),
        )

        row = conn.execute(
            """
            SELECT * FROM draft
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (draft_id, workspace_id, ctx.tenant_id),
        ).fetchone()
        if row is None:
            return None
        record_editorial_event(ctx, conn, "draft", draft_id, "approved", reason=reason)
        return _serialize_draft(row)


def reject_draft(
    ctx: Context,
    conn: sqlite3.Connection,
    draft_id: str,
    *,
    reason: str | None = None,
) -> dict[str, Any] | None:
    workspace_id = ctx.require_scoped_workspace()
    with transaction(conn):
        row = conn.execute(
            """
            SELECT * FROM draft
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (draft_id, workspace_id, ctx.tenant_id),
        ).fetchone()
        if row is None:
            return None

        seal_id = workspace_seal_id(ctx, conn)
        hmac = compute_workspace_hmac_for_row(
            seal_id,
            {
                "id": draft_id,
                "workspace_id": workspace_id,
                "task": row["task"],
                "subject": row["subject"],
                "body_md": row["body_md"],
                "status": "rejected",
                "source_version": row["source_version"],
            },
            "draft",
        )
        conn.execute(
            """
            UPDATE draft
            SET status = 'rejected', updated_at = ?, workspace_hmac = ?, reason = ?
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (utc_now(), hmac, reason, draft_id, workspace_id, ctx.tenant_id),
        )
        row = conn.execute(
            """
            SELECT * FROM draft
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (draft_id, workspace_id, ctx.tenant_id),
        ).fetchone()
        if row is None:
            return None
        record_editorial_event(ctx, conn, "draft", draft_id, "rejected", reason=reason)
        return _serialize_draft(row)


def export_draft(
    ctx: Context,
    conn: sqlite3.Connection,
    draft_id: str,
    *,
    confirmed: bool = False,
) -> dict[str, Any] | None:
    workspace_id = ctx.require_scoped_workspace()
    with transaction(conn):
        row = conn.execute(
            """
            SELECT * FROM draft
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (draft_id, workspace_id, ctx.tenant_id),
        ).fetchone()
        if row is None:
            return None

        data = _serialize_draft(row)
        if data["status"] != "approved":
            raise ValueError("Draft must be approved before export")
        if data["requires_confirmation"] and not confirmed:
            raise ValueError(
                "Draft requires explicit confirmation before export; retry with confirmed=true"
            )

        seal_id = workspace_seal_id(ctx, conn)
        hmac = compute_workspace_hmac_for_row(
            seal_id,
            {
                "id": draft_id,
                "workspace_id": workspace_id,
                "task": row["task"],
                "subject": row["subject"],
                "body_md": row["body_md"],
                "status": "exported",
                "source_version": row["source_version"],
            },
            "draft",
        )
        conn.execute(
            """
            UPDATE draft
            SET status = 'exported', exported_at = ?, updated_at = ?, workspace_hmac = ?
            WHERE id = ? AND workspace_id = ? AND tenant_id = ?
            """,
            (utc_now(), utc_now(), hmac, draft_id, workspace_id, ctx.tenant_id),
        )

    row = conn.execute(
        """
        SELECT * FROM draft
        WHERE id = ? AND workspace_id = ? AND tenant_id = ?
        """,
        (draft_id, workspace_id, ctx.tenant_id),
    ).fetchone()
    return _serialize_draft(row) if row else None
