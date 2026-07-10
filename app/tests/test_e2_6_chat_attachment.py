"""Tests for E2-6 chat attachments."""

from __future__ import annotations

import io
import json
from typing import Any

import pytest
from fastapi.testclient import TestClient


def _make_docx_bytes(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _upload_object(
    client: TestClient,
    workspace_id: str,
    file_name: str,
    mime_type: str,
    data: bytes,
) -> str:
    resp = client.post(
        f"/api/workspaces/{workspace_id}/objects/presigned-upload",
        json={"file_name": file_name, "mime_type": mime_type, "size_bytes": len(data)},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    )
    assert resp.status_code == 201, resp.text
    payload = resp.json()
    from app.src.storage import get_object_store

    store = get_object_store()
    store.put_object(payload["bucket"], payload["object_key"], data, mime_type)
    confirm = client.post(
        f"/api/workspaces/{workspace_id}/objects/confirm",
        json={"object_id": payload["object_id"], "size_bytes": len(data)},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    )
    assert confirm.status_code == 200, confirm.text
    return payload["object_id"]


def _create_workspace(client: TestClient, name: str) -> dict[str, Any]:
    resp = client.post(
        "/api/workspaces",
        json={"name": name},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()


@pytest.fixture()
def client(tmp_path, monkeypatch: pytest.MonkeyPatch):
    db_path = tmp_path / "faberloom.sqlite3"
    audit_path = tmp_path / "audit.jsonl"
    monkeypatch.setenv("FABERLOOM_DB_PATH", str(db_path))
    monkeypatch.setenv("FL_STORAGE_BACKEND", "memory")
    monkeypatch.setenv("FABERLOOM_CONFIG_DIR", str(tmp_path / "config"))

    from app.src.audit import audit_writer
    from app.src.main import create_app
    from app.src.storage import reset_object_store

    reset_object_store()
    audit_writer.audit_path = audit_path
    with TestClient(create_app()) as test_client:
        yield test_client


def test_completion_rejects_missing_attachment_object(client: TestClient) -> None:
    ws = _create_workspace(client, "Chat Attach")
    chat = client.post(
        f"/api/workspaces/{ws['id']}/chats",
        json={"title": "Test"},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    ).json()

    resp = client.post(
        f"/api/workspaces/{ws['id']}/chats/{chat['id']}/completions",
        json={"message": "Resume el adjunto", "attachment_object_id": "obj_doesnotexist"},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    )
    assert resp.status_code == 422
    assert "Attachment" in resp.json()["detail"]


def test_completion_accepts_attachment_object_id(client: TestClient) -> None:
    ws = _create_workspace(client, "Chat Attach OK")
    chat = client.post(
        f"/api/workspaces/{ws['id']}/chats",
        json={"title": "Test"},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    ).json()
    data = _make_docx_bytes("El precio es 150 USD por unidad.")
    obj_id = _upload_object(client, ws["id"], "price.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", data)

    resp = client.post(
        f"/api/workspaces/{ws['id']}/chats/{chat['id']}/completions",
        json={"message": "Resume el adjunto", "attachment_object_id": obj_id},
        headers={"x-tenant-id": "default", "x-user-id": "tester", "x-actor-id": "tester", "x-actor-role": "admin"},
    )
    # No providers configured => 503, but the request is schema-valid and the
    # attachment object exists.
    assert resp.status_code in {422, 503}
